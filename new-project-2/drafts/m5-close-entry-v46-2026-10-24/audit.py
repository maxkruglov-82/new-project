import csv,datetime,hashlib,json,math
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
out=Path(__file__).resolve().parent
root=out.parents[1]
z=json.loads((out/'result.json').read_text())
p=root/'reports/v36-p1-close-2026-10-23/halfyear-candles.json'
raw=json.loads(p.read_text());cs=raw['candles'];ix={c['time']:i for i,c in enumerate(cs)}
ts=lambda s:int(datetime.datetime.fromisoformat(s).timestamp())
end=ts(raw['endExclusive'])
assert z['sourceSha256']==hashlib.sha256((out/'sol_close_entry.pine').read_bytes()).hexdigest()
assert z['inputSha256']==hashlib.sha256(p.read_bytes()).hexdigest()
assert len(cs)==raw['count']==(end-ts(raw['start']))//300
assert all(c['time']==cs[0]['time']+i*300 for i,c in enumerate(cs))
for d in z['days']:
 h=ts(d['referenceClose']);hour=cs[ix[h-3600]:ix[h-3600]+12]
 assert d['high']==max(c['high'] for c in hour) and d['low']==min(c['low'] for c in hour)
 assert datetime.date.fromisoformat(d['date']).weekday()<5
for e in z['p1Audit']:
 j=ix[ts(e['pivotAt'])];k=ix[ts(e['priorCloseAt'])];i=ix[ts(e['confirmedAt'])];assert k<=j<i
 ref=ts(e['pivotAt'][:10]+'T15:00:00+07:00');assert ref<=cs[k]['time']<=cs[j]['time']
 hour=cs[ix[ref-3600]:ix[ref-3600]+12];hi=max(b['high'] for b in hour);lo=min(b['low'] for b in hour)
 assert e['h1High']==hi and e['h1Low']==lo
 correct=lambda b:b['close']<lo if e['side']==-1 else b['close']>hi
 assert next(n for n in range(ix[ref],j+1) if correct(cs[n]))==k
 assert e['priorClose']==cs[k]['close'] and e['pivotClose']==cs[j]['close']
for t in z['trades']:
 i=ix[ts(t['entryBarStart'])];b=cs[i];assert b['low']<=t['entry']<=b['high']
 assert datetime.datetime.fromisoformat(t['entryBarStart']).weekday()<5
 assert t['entryBarStart'][:10]==t['referenceDate'] and t['entryBarStart'][11:16]>='15:00'
 assert t['entryAt'][:10]==t['referenceDate'] and ts(t['entryAt'])==b['time']+300
 assert t['entry']==b['close']
 assert (b['high']<t['stop'] if t['side']=='SHORT' else b['low']>t['stop'])
 assert (t['entry']>=t['rrThreshold'] if t['side']=='SHORT' else t['entry']<=t['rrThreshold'])
 assert math.isclose(t['rrThreshold'],(t['stop']+t['swingTarget'])/2,abs_tol=1e-9)
 assert math.isclose(t['stop'],t['fiboZero']+1.25*(t['fiboOne']-t['fiboZero'])+(.10 if t['side']=='SHORT' else -.10),abs_tol=1e-9)
 assert math.isclose(t['target'],t['swingTarget'],abs_tol=1e-9)
 assert t['rewardRisk']>=1-1e-8
 j=int(t['setupConfirmedBar']);assert j<i
 assert cs[j]['close']==t['setupBreakClose']
 assert (cs[j]['close']<t['setupBreakLevel']) if t['side']=='SHORT' else (cs[j]['close']>t['setupBreakLevel'])
 expiry=ts(t['referenceDate']+'T15:00:00+07:00')+86400
 status='unresolved_at_end' if end<expiry else 'unresolved_at_new_h1';exit_time=None
 for b in cs[i+1:]:
  if b['time']>=min(expiry,end):break
  sl=b['high']>=t['stop'] if t['side']=='SHORT' else b['low']<=t['stop']
  tp=b['low']<=t['target'] if t['side']=='SHORT' else b['high']>=t['target']
  if sl or tp:status='ambiguous' if sl and tp else 'stop' if sl else 'target';exit_time=b['time'];break
 assert status==t['status'],(t['id'],status,t['status'])
 if exit_time is not None:assert ts(t['exitBarStart'])==exit_time
 if status=='target':assert math.isclose(t['r'],abs(t['target']-t['entry'])/abs(t['stop']-t['entry']))
 if status=='stop':assert t['r']==-1
for d in z['days']:
 own=[t for t in z['trades'] if t['referenceDate']==d['date']]
 for i,t in enumerate(own):
  if t['status']=='target':assert i==len(own)-1

for day in z['days']:
 own=[t for t in z['trades'] if t['referenceDate']==day['date']]
 assert len(own)<=2
 if len(own)==2:
  assert own[0]['status']=='stop' and own[0]['side']!=own[1]['side']
  assert ts(own[1]['entryBarStart'])>=ts(own[0]['exitBarEnd'])
for e in z['failedReversalAudit']:
 j=ix[ts(e['sweepAt'])];i=ix[ts(e['at'])];assert j<i
 if e['side']==1:
  assert cs[j]['low']==e['original1'] and cs[j]['low']<e['sweptLevel']<=cs[j]['close'] and cs[i]['close']>e['original2']
 else:
  assert cs[j]['high']==e['original1'] and cs[j]['high']>e['sweptLevel']>=cs[j]['close'] and cs[i]['close']<e['original2']
for day in ['2026-10-16','2026-10-19','2026-10-20','2026-10-22']:
 control=json.loads((out/(day+'-control.json')).read_text());assert control['sourceSha256']==z['sourceSha256']
 keys=['side','entryBarStart','entry','stop','target','status','exitBarStart']
 assert [{k:t[k] for k in keys} for t in z['trades'] if t['referenceDate']==day]==[{k:t[k] for k in keys} for t in control['trades']]
months=z['monthly']; perf=z['performance']
for key in ['entries','closed','wins','stops','unresolved','ambiguous']:
 assert sum(m[key] for m in months)==perf[key]
assert math.isclose(sum(m['netR'] for m in months),perf['netR'])
wb=Workbook();ws=wb.active;ws.title='2026 помесячно'
keys=['month','entries','closed','wins','stops','winRate','netR','unresolved','ambiguous'];ws.append(keys)
worst={m['month'] for m in sorted(months,key=lambda m:(m['winRate'],m['netR']))[:3]}
for m in months:
 ws.append([m[k] for k in keys])
 if m['month'] in worst:
  for cell in ws[ws.max_row]:cell.font=Font(color='9C0006');cell.fill=PatternFill('solid',fgColor='FFC7CE')
ws.append(['ИТОГО']+[perf[k] for k in keys[1:]])
for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=19
ws.freeze_panes='A2';wb.save(out/'monthly.xlsx')
lines=['# Solana M5 v46 — черновик, полгода — 2026 помесячно','', 'MEXC SOL_USDT, 23 апреля — 22 октября 2026, UTC+7. Точный опубликованный Pine. Без комиссий, проскальзывания и funding. Незавершённые и неоднозначные сделки исключены из винрейта.','', '| Месяц | Закрытые | TP | SL | Винрейт | R |','|---|---:|---:|---:|---:|---:|']
for m in months:lines.append(f"| {m['month']} | {m['closed']} | {m['wins']} | {m['stops']} | {m['winRate']:.2f}% | {m['netR']:+.2f} |")
lines+=['',str(perf),'','Три худших месяца по винрейту выделены красным в monthly.xlsx.']
(out/'report.md').write_text('\n'.join(lines)+'\n')
audit={'status':'PASS','sourceSha256':z['sourceSha256'],'inputSha256':z['inputSha256'],'candles':len(cs),'trades':len(z['trades']),'p1Events':len(z['p1Audit']),'checks':['Continuous candles','H1 ranges','P1 prior closes','Entry price/time','Stop1.25 based on P3/P2, TP=P3, RR1 threshold midpoint SL/TP; execution at close; RR>=1','First subsequent M5 exit','Monthly totals','Max2 trades/day; second opposite after first SL; before midnight; confirmed close','Failed-break reversals verified from raw candles','Oct19/20/22 ledgers match standalone controls']}
(out/'audit.json').write_text(json.dumps(audit,indent=2)+'\n');print(json.dumps(perf,indent=2));print(json.dumps(months,indent=2))
