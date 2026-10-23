import csv,datetime,hashlib,json,math
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
out=Path(__file__).resolve().parent
root=out.parents[1]
z=json.loads((out/'halfyear-v39.json').read_text())
a=json.loads((out/'halfyear-comparison.json').read_text())
p=root/'reports/v36-p1-close-2026-10-23/halfyear-candles.json'
raw=json.loads(p.read_text());cs=raw['candles'];ix={c['time']:i for i,c in enumerate(cs)}
ts=lambda s:int(datetime.datetime.fromisoformat(s).timestamp())
end=ts(raw['endExclusive'])
assert z['sourceSha256']==hashlib.sha256((out/'sol_equal_p3.pine').read_bytes()).hexdigest()
assert z['inputSha256']==hashlib.sha256(p.read_bytes()).hexdigest()
assert len(cs)==raw['count']==(end-ts(raw['start']))//300
assert all(c['time']==cs[0]['time']+i*300 for i,c in enumerate(cs))
for d in z['days']:
 h=ts(d['referenceClose']);hour=cs[ix[h-3600]:ix[h-3600]+12]
 assert d['high']==max(c['high'] for c in hour) and d['low']==min(c['low'] for c in hour)
 assert datetime.date.fromisoformat(d['date']).weekday()<5
for e in z['p1Audit']:
 j=ix[ts(e['pivotAt'])];k=ix[ts(e['priorCloseAt'])];i=ix[ts(e['confirmedAt'])];assert k<=j<i
 ref=ts(e['pivotAt'][:10]+'T15:00:00+07:00');assert ref<=cs[k]['time']<=cs[j]['time']
 hour=cs[ix[ref-3600]:ix[ref-3600]+12];hi=max(b['high'] for b in hour);lo=min(b['low'] for b in hour)
 assert e['h1High']==hi and e['h1Low']==lo
 correct=lambda b:b['close']<lo if e['side']==-1 else b['close']>hi
 assert next(n for n in range(ix[ref],j+1) if correct(cs[n]))==k
 assert e['priorClose']==cs[k]['close'] and e['pivotClose']==cs[j]['close']
for t in z['trades']:
 i=ix[ts(t['entryBarStart'])];b=cs[i];assert b['low']<=t['entry']<=b['high']
 assert datetime.datetime.fromisoformat(t['entryBarStart']).weekday()<5
 assert t['entryBarStart'][:10]==t['referenceDate'] and t['entryBarStart'][11:16]>='15:00'
 assert (b['open']<=t['entry'] if t['side']=='SHORT' else b['open']>=t['entry'])
 assert math.isclose(t['entry'],(t['fiboStop']+t['swingTarget'])/2,abs_tol=1e-9)
 assert math.isclose(t['stop']-t['fiboStop'],.03 if t['side']=='SHORT' else -.03,abs_tol=1e-9)
 assert math.isclose(t['target']-t['swingTarget'],.10 if t['side']=='SHORT' else -.10,abs_tol=1e-9)
 expiry=ts(t['referenceDate']+'T15:00:00+07:00')+86400
 status='unresolved_at_end' if end<expiry else 'unresolved_at_new_h1';exit_time=None
 for b in cs[i+1:]:
  if b['time']>=min(expiry,end):break
  sl=b['high']>=t['stop'] if t['side']=='SHORT' else b['low']<=t['stop']
  tp=b['low']<=t['target'] if t['side']=='SHORT' else b['high']>=t['target']
  if sl or tp:status='ambiguous' if sl and tp else 'stop' if sl else 'target';exit_time=b['time'];break
 assert status==t['status'],(t['id'],status,t['status'])
 if exit_time is not None:assert ts(t['exitBarStart'])==exit_time
 if status=='target':assert math.isclose(t['r'],abs(t['target']-t['entry'])/abs(t['stop']-t['entry']))
 if status=='stop':assert t['r']==-1
for d in z['days']:
 own=[t for t in z['trades'] if t['referenceDate']==d['date']]
 for i,t in enumerate(own):
  if t['status']=='target':assert i==len(own)-1
s=z['summary'];months=a['monthly']
for key,summarykey in [('entered','enteredTrades'),('closed','totalTrades'),('wins','targets'),('losses','stops'),('unresolved','unresolved'),('ambiguous','ambiguous')]:
 assert sum(m['v39'][key] for m in months)==s[summarykey]
assert math.isclose(sum(m['v39']['netR'] for m in months),s['netRKnownBeforeCosts'])
assert math.isclose(s['winRateKnown'],s['targets']/(s['targets']+s['stops']))
for name,n in [('halfyear-monthly.csv',len(months)),('halfyear-trades.csv',len(z['trades']))]:
 with (out/name).open(encoding='utf-8-sig') as f:assert len(list(csv.DictReader(f)))==n
# Keep the special M1 example separate from the M5-only backtest.
t=next(t for t in z['trades'] if t['referenceDate']=='2026-10-22' and t['side']=='SHORT')
assert t['entryBarStart']=='2026-10-22T16:30:00.000+07:00' and t['status']=='stop'
assert t['exitBarStart']=='2026-10-22T16:45:00.000+07:00'
wb=Workbook();sheet=wb.active;sheet.title='Сравнение'
rows=[['SOL/USDT MEXC M5 • 23.04–22.10.2026 • UTC+7'],['Вариант','Всего входов','Закрытых','Тейки','Стопы','Незавершённые','Неоднозначные','Винрейт','Результат R','Макс. просадка R','Профит-фактор']]
for version in ['v38','v39']:
 q=a[version];rows.append([version]+[q[k] for k in ['entered','closed','wins','losses','unresolved','ambiguous','winRate','netR','maxDrawdownR','profitFactor']])
for row in rows:sheet.append(row)
for r in [3,4]:sheet.cell(r,8).number_format='0.00%'
monthly=wb.create_sheet('Помесячно v39')
monthly.append(['Месяц','Всего входов','Закрытых','Тейки','Стопы','Незавершённые','Неоднозначные','Винрейт v39','R v39','Винрейт v38','R v38'])
for m in months:
 q=m['v39'];monthly.append([m['month']]+[q[k] for k in ['entered','closed','wins','losses','unresolved','ambiguous','winRate','netR']]+[m['v38']['winRate'],m['v38']['netR']])
 for col in [8,10]:monthly.cell(monthly.max_row,col).number_format='0.00%'
trades=wb.create_sheet('Сделки v39');keys=list(z['trades'][0]);trades.append(keys)
for t in z['trades']:trades.append([t.get(k) for k in keys])
for ws,header in [(sheet,2),(monthly,1),(trades,1)]:
 ws.freeze_panes=f'A{header+1}';ws.auto_filter.ref=f'A{header}:{ws.cell(ws.max_row,ws.max_column).coordinate}'
 for cell in ws[header]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17365D')
 for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(40,max(14,max(len(str(c.value or '')) for c in col)+2))
 for row in ws.iter_rows(min_row=header+1):
  for cell in row:
   if isinstance(cell.value,float) and cell.number_format=='General':cell.number_format='0.0000'
wb.save(out/'halfyear-comparison.xlsx')
check=load_workbook(out/'halfyear-comparison.xlsx',data_only=True)
assert check['Сравнение']['B4'].value==s['enteredTrades'] and check['Сделки v39'].max_row==len(z['trades'])+1
labels=['','Январь','Февраль','Март','Апрель 23–30','Май','Июнь','Июль','Август','Сентябрь','Октябрь 1–22']
pct=lambda x:f'{x*100:.2f}%'
lines=['# Solana M5 v39 — сравнение с v38 за полгода','','23 апреля — 22 октября 2026 включительно, UTC+7. MEXC SOL_USDT perpetual futures; 52 536 непрерывных M5. v39 допускает равный P1 экстремум как P3 и следующий отдельный P2 слабее пропущенного соседнего свинга. Остальные правила и буферы сохранены. Выполнен фактический Pine-исходник, без переписывания стратегии. v38 — ранее проверенный журнал на тех же данных, проверены SHA256 и эквивалентность опубликованного исходника.','','| Версия | Входы | Закрытые | Тейки | Стопы | Винрейт | R |','|---|---:|---:|---:|---:|---:|---:|']
for version in ['v38','v39']:
 q=a[version];lines.append(f"| {version} | {q['entered']} | {q['closed']} | {q['wins']} | {q['losses']} | {pct(q['winRate'])} | {q['netR']:+.4f} |")
lines+=['','| Месяц | Входы | Закрытые | Тейки | Стопы | Винрейт v39 | R v39 |','|---|---:|---:|---:|---:|---:|---:|']
for m in months:
 q=m['v39'];lines.append(f"| {labels[int(m['month'][5:])]} | {q['entered']} | {q['closed']} | {q['wins']} | {q['losses']} | {pct(q['winRate'])} | {q['netR']:+.4f} |")
lines+=['',f"Даты с изменёнными сделками: {', '.join(d['date'] for d in a['changedDays'])}.",'','Винрейт = тейки / (тейки + стопы). Незавершённые и неоднозначные исходы исключены из винрейта и результата R; количества приведены в Excel и JSON. Результат R до комиссий, финансирования и проскальзывания. Из-за тейка на 0,10 USDT перед целью и стопа на 0,03 USDT за Фибо фактический RR ниже 1:1.','', 'Исторический вход проверяется по последующей свече ретеста, выход — начиная со свечи после входа. Порядок касаний внутри M5 неизвестен; одновременный SL/TP помечается неоднозначным. Новый опорный час снимает незавершённую позицию без фиктивного результата. После тейка новых входов в этот опорный день нет.','','22 октября: M5-only SHORT16:30,стоп16:45. Более подробное воспроизведение по M1 даёт иной ранний вход и тейк; он не подставлен в этот M5-бэктест. Черновик не опубликован.','','[Excel](halfyear-comparison.xlsx) · [Журнал CSV](halfyear-trades.csv) · [Сравнение JSON](halfyear-comparison.json) · [Аудит](halfyear-audit.json)']
(out/'halfyear-report.md').write_text('\n'.join(lines)+'\n')
audit={'status':'PASS','candles':len(cs),'verifiedReferences':len(z['days']),'verifiedP1Events':len(z['p1Audit']),'verifiedTrades':len(z['trades']),'sourceSha256':z['sourceSha256'],'inputSha256':z['inputSha256'],'checks':['Continuous M5 data and actual reference H1','Each P1 prior close independently reconstructed, without future close','Entry price touched, time and weekdays','Unchanged E/SL/TP formulas','First subsequent OHLC exit independently reconstructed','No new entry after TP','Monthly and CSV reconciliation','Oct22 historical M5 loss retained; separate minute win not substituted','Workbook values checked']}
(out/'halfyear-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(audit,ensure_ascii=False,indent=2))
