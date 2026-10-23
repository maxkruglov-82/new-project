import csv,datetime,hashlib,json,math
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
out=Path(__file__).resolve().parents[1]/'reports/v40-ytd-2026-10-24'
root=out.parents[1]
z=json.loads((out/'result.json').read_text())
p=root/'reports/mexc-sol-usdt-m5-2026-01-01_2026-10-23.json'
raw=json.loads(p.read_text());cs=raw['candles'];ix={c['time']:i for i,c in enumerate(cs)}
ts=lambda s:int(datetime.datetime.fromisoformat(s).timestamp())
end=ts(raw['endExclusive'])
assert z['sourceSha256']==hashlib.sha256((root/'sol_short_setup.pine').read_bytes()).hexdigest()
assert z['inputSha256']==hashlib.sha256(p.read_bytes()).hexdigest()
assert len(cs)==raw['count']==(end-ts(raw['start']))//300
assert all(c['time']==cs[0]['time']+i*300 for i,c in enumerate(cs))
for d in z['days']:
 h=ts(d['referenceClose']);hour=cs[ix[h-3600]:ix[h-3600]+12]
 assert d['high']==max(c['high'] for c in hour) and d['low']==min(c['low'] for c in hour)
 assert datetime.date.fromisoformat(d['date']).weekday()<5
for e in z['p1Audit']:
 j=ix[ts(e['pivotAt'])];k=ix[ts(e['priorCloseAt'])];i=ix[ts(e['confirmedAt'])];assert k<=j<i
 ref=ts(e['pivotAt'][:10]+'T15:00:00+07:00');assert ref<=cs[k]['time']<=cs[j]['time']
 hour=cs[ix[ref-3600]:ix[ref-3600]+12];hi=max(b['high'] for b in hour);lo=min(b['low'] for b in hour)
 assert e['h1High']==hi and e['h1Low']==lo
 correct=lambda b:b['close']<lo if e['side']==-1 else b['close']>hi
 assert next(n for n in range(ix[ref],j+1) if correct(cs[n]))==k
 assert e['priorClose']==cs[k]['close'] and e['pivotClose']==cs[j]['close']
for t in z['trades']:
 i=ix[ts(t['entryBarStart'])];b=cs[i];assert b['low']<=t['entry']<=b['high']
 assert datetime.datetime.fromisoformat(t['entryBarStart']).weekday()<5
 assert t['entryBarStart'][:10]==t['referenceDate'] and t['entryBarStart'][11:16]>='15:00'
 assert (b['open']<=t['entry'] if t['side']=='SHORT' else b['open']>=t['entry'])
 assert math.isclose(t['entry'],(t['fiboStop']+t['swingTarget'])/2,abs_tol=1e-9)
 assert math.isclose(t['stop']-t['fiboStop'],.03 if t['side']=='SHORT' else -.03,abs_tol=1e-9)
 assert math.isclose(t['target']-t['swingTarget'],.10 if t['side']=='SHORT' else -.10,abs_tol=1e-9)
 expiry=ts(t['referenceDate']+'T15:00:00+07:00')+86400
 status='unresolved_at_end' if end<expiry else 'unresolved_at_new_h1';exit_time=None
 for b in cs[i+1:]:
  if b['time']>=min(expiry,end):break
  sl=b['high']>=t['stop'] if t['side']=='SHORT' else b['low']<=t['stop']
  tp=b['low']<=t['target'] if t['side']=='SHORT' else b['high']>=t['target']
  if sl or tp:status='ambiguous' if sl and tp else 'stop' if sl else 'target';exit_time=b['time'];break
 assert status==t['status'],(t['id'],status,t['status'])
 if exit_time is not None:assert ts(t['exitBarStart'])==exit_time
 if status=='target':assert math.isclose(t['r'],abs(t['target']-t['entry'])/abs(t['stop']-t['entry']))
 if status=='stop':assert t['r']==-1
for d in z['days']:
 own=[t for t in z['trades'] if t['referenceDate']==d['date']]
 for i,t in enumerate(own):
  if t['status']=='target':assert i==len(own)-1

months=z['monthly']; perf=z['performance']
for key in ['entries','closed','wins','losses','unresolved','ambiguous']:
 assert sum(m[key] for m in months)==perf[key]
assert math.isclose(sum(m['netR'] for m in months),perf['netR'])
wb=Workbook();ws=wb.active;ws.title='2026 помесячно'
keys=['month','entries','closed','wins','losses','winRate','netR','unresolved','ambiguous'];ws.append(keys)
worst={m['month'] for m in sorted(months,key=lambda m:(m['winRate'],m['netR']))[:3]}
for m in months:
 ws.append([m[k] for k in keys])
 if m['month'] in worst:
  for cell in ws[ws.max_row]:cell.font=Font(color='9C0006');cell.fill=PatternFill('solid',fgColor='FFC7CE')
ws.append(['ИТОГО']+[perf[k] for k in keys[1:]])
for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=19
ws.freeze_panes='A2';wb.save(out/'monthly.xlsx')
lines=['# Solana M5 v40 — 2026 помесячно','', 'MEXC SOL_USDT, 1 января — 22 октября 2026, UTC+7. Точный опубликованный Pine. Без комиссий, проскальзывания и funding. Незавершённые и неоднозначные сделки исключены из винрейта.','', '| Месяц | Закрытые | TP | SL | Винрейт | R |','|---|---:|---:|---:|---:|---:|']
for m in months:lines.append(f"| {m['month']} | {m['closed']} | {m['wins']} | {m['losses']} | {m['winRate']:.2f}% | {m['netR']:+.2f} |")
lines+=['',str(perf),'','Три худших месяца по винрейту выделены красным в monthly.xlsx.']
(out/'report.md').write_text('\n'.join(lines)+'\n')
audit={'status':'PASS','sourceSha256':z['sourceSha256'],'inputSha256':z['inputSha256'],'candles':len(cs),'trades':len(z['trades']),'p1Events':len(z['p1Audit']),'checks':['Continuous candles','H1 ranges','P1 prior closes','Entry price/time','SL/TP formulas','First subsequent M5 exit','Monthly totals','Exact halfyear ledger match in runner']}
(out/'audit.json').write_text(json.dumps(audit,indent=2)+'\n');print(json.dumps(perf,indent=2));print(json.dumps(months,indent=2))
